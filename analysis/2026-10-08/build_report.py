"""Build the review report from its Markdown master using python-docx."""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

O=Path(__file__).resolve().parent
src=O/'Bao_cao_danh_gia_huong_trien_khai.md'
text=src.read_text(); lines=text.splitlines()
refs={}
for line in lines:
 m=re.match(r'- \[(S\d+)\].*?(https://\S+)',line)
 if m: refs[m[1]]=m[2].rstrip('.,;')
(O/'source_register.json').write_text(json.dumps(refs,ensure_ascii=False,indent=2))
d=Document(); sec=d.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.75);sec.footer_distance=Cm(.75)
styles=d.styles
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','Caption']:
 s=styles[n];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial')
 s.element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Arial')
 s.paragraph_format.widow_control=True
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.line_spacing=1.12
styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Title'].font.size=Pt(23);styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(10)
styles['Subtitle'].font.size=Pt(10)
for n,size in [('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
 s=styles[n];s.font.size=Pt(size);s.font.bold=True
 s.paragraph_format.space_before=Pt(14 if n=='Heading 1' else 9)
 s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
styles['List Bullet'].paragraph_format.space_after=Pt(4)
styles['List Bullet'].paragraph_format.line_spacing=1.08

head=sec.header.paragraphs[0];head.text='DSP391m  |  Group 8  |  Đánh giá hướng triển khai'
head.style=styles['Normal'];head.paragraph_format.space_after=Pt(0)
for r in head.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('08/10/2026  ·  ');r.font.size=Pt(8)
for code in ['PAGE','NUMPAGES']:
 if code=='NUMPAGES':foot.add_run(' / ').font.size=Pt(8)
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),code)
 rr=OxmlElement('w:r');rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'16');rp.append(sz);rr.append(rp)
 tt=OxmlElement('w:t');tt.text='1';rr.append(tt);fld.append(rr);foot._p.append(fld)

def link(p,label,url,size=None):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
 color=OxmlElement('w:color');color.set(qn('w:val'),'244A70');rp.append(color)
 if size:
  z=OxmlElement('w:sz');z.set(qn('w:val'),str(int(size*2)));rp.append(z)
 r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

pattern=re.compile(r'(\*\*.*?\*\*|`[^`]+`|https://[^\s]+|\[S\d+(?:, S\d+)*\])')
def inline(p,s,size=None):
 for part in pattern.split(s):
  if not part:continue
  if part.startswith('https://'):
   clean=part.rstrip('.,;');link(p,clean,clean,size)
   if part!=clean:p.add_run(part[len(clean):])
  elif re.fullmatch(r'\[S\d+(?:, S\d+)*\]',part):
   ids=part[1:-1].split(', ');p.add_run('[')
   for idx,k in enumerate(ids):
    if idx:p.add_run(', ')
    if k in refs:link(p,k,refs[k],size)
    else:p.add_run(k)
   p.add_run(']')
  else:
   bold=part.startswith('**');code=part.startswith('`')
   r=p.add_run(part[2:-2] if bold else part[1:-1] if code else part)
   r.bold=bold
   if code:r.font.name='Consolas';r.font.size=Pt(size or 9)
   elif size:r.font.size=Pt(size)

def table(block):
 rows=[[x.strip() for x in l.strip().strip('|').split('|')] for l in block]
 rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',x) for x in r)]
 n=len(rows[0]);tb=d.add_table(rows=1,cols=n);tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
 if n==2: widths=[7.5,9.5]
 elif n==3: widths=[4.2,5.5,7.3]
 elif n==4: widths=[3.5,5.0,5.5,3.0]
 elif n==5: widths=[3.1,2.4,5.0,3.25,3.25]
 else:widths=[17/n]*n
 # Wider narrative columns for specific tables.
 if rows[0][0]=='Nguồn':widths=[2.5,5.2,4.5,4.8]
 if rows[0][0]=='Phương án':widths=[3.0,5.0,5.5,3.5]
 if rows[0][0]=='Số thuốc k':widths=[2.0,3.6,4.4,3.7,3.3]
 if rows[0][0]=='Phạm vi':widths=[3.8,2.3,5.2,2.85,2.85]
 for col,w in zip(tb.columns,widths):col.width=Cm(w)
 pr=tb._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for edge in ('top','left','bottom','right','insideH','insideV'):
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 margins=OxmlElement('w:tblCellMar')
 for side,val in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
  e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
 pr.append(margins)
 for i,row in enumerate(rows):
  cells=tb.rows[0].cells if i==0 else tb.add_row().cells
  trpr=tb.rows[i]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,(cell,s) in enumerate(zip(cells,row)):
   cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'23455E' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shade)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.06
   p.paragraph_format.keep_with_next=i==0
   inline(p,s,9.3)
   if i==0:
    for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
   elif re.fullmatch(r'[\d.,%– /]+',s):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 # separation avoids visually touching the next paragraph
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.add_run().font.size=Pt(3)

i=0;reference_section=False
while i<len(lines):
 l=lines[i].strip()
 if not l:i+=1;continue
 if l.startswith('|'):
  block=[]
  while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
  table(block);continue
 if l.startswith('# '):
  p=d.add_paragraph(l[2:],'Title')
 elif l.startswith('## '):
  title=l[3:]
  if title.startswith('Phụ lục B'):reference_section=True
  p=d.add_paragraph(title,'Heading 1')
 elif l.startswith('### '):p=d.add_paragraph(l[4:],'Heading 2')
 elif l.startswith('- '):
  p=d.add_paragraph(style='Normal' if reference_section else 'List Bullet');inline(p,l[2:],9.3 if reference_section else None)
  if reference_section:
   p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.05
 else:
  p=d.add_paragraph();inline(p,l)
 i+=1

# Explicitly strip inherited title borders and theme colors.
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=styles[name]
 for c in s.element.findall('.//'+qn('w:color')):
  for a in ('themeColor','themeShade','themeTint'):c.attrib.pop(qn('w:'+a),None)
  c.set(qn('w:val'),'000000')
 for b in s.element.findall('.//'+qn('w:pBdr')):b.getparent().remove(b)

d.core_properties.title='Đánh giá tính khả thi và hướng triển khai dự án tương tác đa thuốc'
d.core_properties.subject='Kiểm toán dữ liệu và kế hoạch nghiên cứu DSP391m'
d.core_properties.author='Báo cáo nghiên cứu cho Group 8'
d.core_properties.keywords='Decagon, HODDI, DDI, Graph RAG, feasibility'
out=O/'Bao_cao_danh_gia_huong_trien_khai.docx';d.save(out)
print(out)
print('Tables:',len(d.tables),'Paragraphs:',len(d.paragraphs),'Sources:',len(refs))
